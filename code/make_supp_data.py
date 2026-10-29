#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Supplementary_Data.xlsx：**底層資料**，不是圖上的值。

審稿人（AoM #42 Reviewer 1，2026-10-28）：「While figures and tables are provided,
the underlying data (e.g. genome accessions, coordinates of BGCs, antibiotic
resistance genes) are not provided as spreadsheets which should be standard.」

Source_Data.xlsx 是「一格圖一個分頁」，放的是畫出來的那些數字；這一本放的是
那些數字底下的逐列資料：每一株的 accession 與 metadata、562 個基因簇的座標、
簇窗內每一個 CDS、每一個抗藥基因呼叫與它的座標、每一簇到最近抗藥基因的距離、
定序型、引子與預測質量。兩本分開，不要合併：讀者要的是「表在哪」而不是
「從 53 個面板分頁裡自己湊」。

  python3 code/make_supp_data.py   ->  out/Supplementary_Data.xlsx
"""
import csv, json, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(HERE, "out")
WB = Workbook()
WB.remove(WB.active)
INDEX = []


def rows_of(name, delim=None):
    """讀 out/<name>，回傳 (header, rows)。分隔字元由副檔名決定。"""
    path = os.path.join(OUT, name)
    if not os.path.exists(path):
        return None, None
    if delim is None:
        delim = "\t" if name.endswith(".tsv") else ","
    with open(path, newline="") as fh:
        r = list(csv.reader(fh, delimiter=delim))
    return (r[0], r[1:]) if r else (None, None)


def sheet(tag, what, source, header, rows, note=""):
    ws = WB.create_sheet(tag[:31])
    ws.cell(1, 1, f"{tag} — {what}").font = Font(bold=True)
    ws.cell(2, 1, f"out/{source}" + (f" · {note}" if note else "")).font = \
        Font(italic=True, size=9)
    for j, h in enumerate(header, start=1):
        c = ws.cell(4, j, h)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    for i, r in enumerate(rows, start=5):
        for j, v in enumerate(r, start=1):
            # 數字留成數字，讀者才能直接排序與計算
            try:
                ws.cell(i, j, int(v) if str(v).lstrip("-").isdigit() else float(v))
            except (TypeError, ValueError):
                ws.cell(i, j, v)
    for j in range(1, len(header) + 1):
        widths = [len(str(header[j - 1]))]
        widths += [len(str(r[j - 1])) for r in rows[:60] if len(r) >= j]
        ws.column_dimensions[get_column_letter(j)].width = \
            max(10, min(44, max(widths) + 2))
    ws.freeze_panes = "A5"
    INDEX.append([tag, what, f"out/{source}", len(rows)])
    return ws


# tag, 來源檔, 說明, 備註
SPEC = [
    ("D1_Genomes", "metadata.csv",
     "every genome in the panel: assembly accession, species, strain, "
     "BioProject, BioSample, submitter, dates, host and isolation source, "
     "assembly size, GC, replicon and gene counts",
     "the inclusion list; RefSeq complete genomes retrieved 20 September 2026"),
    ("D2_Clusters", "clusters.tsv",
     "every biosynthetic gene cluster called in the panel, with its replicon, "
     "coordinates, span, scanning window, anchor enzymes and domain architecture",
     "one row per cluster"),
    ("D3_Cluster_CDS", "cluster_cds.tsv",
     "every coding sequence inside every cluster window, with coordinates, "
     "strand, protein accession, RefSeq product and the domains matched",
     "the gene-level coordinates behind D2"),
    ("D4_AMR_calls", "amr_located.csv",
     "every acquired resistance determinant called, with the protein that "
     "carries it, its replicon and its coordinates",
     ""),
    ("D5_Cluster_to_AMR", "colocalisation.csv",
     "per cluster: the acquired determinants on the same replicon and the "
     "distance to the nearest one",
     "the distances reported in the text"),
    ("D6_Distance_null", "distance_null.csv",
     "per cluster: observed distance to the nearest determinant on the same "
     "replicon against the median of the permutation null",
     ""),
    ("D7_Families", "families.tsv",
     "the cluster families: size, species spread, replicon location, span, "
     "domain content and the RefSeq products of their members",
     ""),
    ("D8_Genome_matrix", "genome_matrix.csv",
     "genome x family carriage matrix, with the metadata columns used for "
     "stratification",
     ""),
    ("D9_MLST", "mlst.csv",
     "seven-locus sequence types and the allele calls behind them",
     "untyped rows are genomes whose alleles are newer than the public "
     "database access cap"),
    ("D10_Primers", "primers.csv",
     "the PCR scheme: primer sequences, predicted product, melting "
     "temperatures and in-silico performance",
     "in-silico only; no primer was tested at the bench"),
    ("D11_External_PCR", "external_pcr.csv",
     "in-silico PCR on the external sets: the draft assemblies and the "
     "non-target species",
     ""),
    ("D12_Precursor_mass", "precursor_masses.csv",
     "predicted mature core peptides and their mass ladders",
     "predictions only; no peptide was expressed or analysed by mass "
     "spectrometry"),
    ("D13_antiSMASH", "antismash_compare.csv",
     "atlas clusters against the antiSMASH regions of the same genomes",
     "written by code/s55_antismash_compare.py"),
    ("D14_AMR_containment", "amr_per_genome_containment.csv",
     "the earlier k-mer containment calls, kept for the comparison with "
     "AMRFinderPlus",
     "superseded; retained so the two call sets can be compared"),
    ("D15_antiSMASH_regions", "antismash_regions.csv",
     "every region antiSMASH called in the panel, with its product, whether "
     "that product falls within the atlas's profile scope, and whether an "
     "atlas cluster overlaps it",
     "antiSMASH 8.0.4 on the same RefSeq gene calls; written by "
     "code/s54_antismash.py"),
    ("D16_MIBiG", "mibig_hits.csv",
     "the best MIBiG match of every antiSMASH region, by knownclusterblast, "
     "with the atlas family of the cluster that overlaps it where there is one",
     "similarity is the share of the reference cluster's genes found in the "
     "region, not sequence identity; MIBiG 4.0"),
    ("D17_Pseudogenes", "pseudogenes.csv",
     "anchor genes of a family that RefSeq annotates as pseudogenes, with the "
     "reason given and whether the atlas counts that genome as a carrier",
     "these are invisible to a proteome-based call; written by "
     "code/s57_pseudogenes.py"),
]


def main():
    for tag, src, what, note in SPEC:
        header, rows = rows_of(src)
        if header is None:
            print(f"   skip {tag}: out/{src} not present", file=sys.stderr)
            continue
        sheet(tag, what, src, header, rows, note)

    idx = WB.create_sheet("Index")
    idx.cell(1, 1, "Supplementary data: the underlying tables").font = \
        Font(bold=True, size=12)
    idx.cell(2, 1, "One sheet per underlying table. Figure-panel values are in "
                   "Source_Data.xlsx instead; this workbook holds the rows those "
                   "values were computed from.").font = Font(italic=True, size=9)
    for j, h in enumerate(["sheet", "what it holds", "file", "rows"], start=1):
        idx.cell(4, j, h).font = Font(bold=True)
    for i, r in enumerate(INDEX, start=5):
        for j, v in enumerate(r, start=1):
            idx.cell(i, j, v).alignment = Alignment(vertical="top",
                                                    wrap_text=(j == 2))
    idx.column_dimensions["A"].width = 20
    idx.column_dimensions["B"].width = 80
    idx.column_dimensions["C"].width = 38
    idx.column_dimensions["D"].width = 8
    WB._sheets.insert(0, WB._sheets.pop(WB._sheets.index(idx)))

    dest = os.path.join(OUT, "Supplementary_Data.xlsx")
    WB.save(dest)
    print("->", dest, f"({len(WB.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
